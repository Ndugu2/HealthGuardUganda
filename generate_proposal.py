import sys
import subprocess

try:
    import docx
except ImportError:
    subprocess.check_call([sys.executable, "-m", "pip", "install", "python-docx"])
    import docx

from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()

# Title Page
doc.add_heading('HealthGuard Uganda', 0).alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.add_paragraph('Systemic Health Misinformation Screening & Offline-First Community Health Platform').alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.add_paragraph('A proposal for a resilient offline-first mobile/web health information, triage, and infodemic-management platform serving rural and semi-urban Uganda').alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.add_page_break()

# Executive Summary
doc.add_heading('Executive Summary', level=1)
doc.add_paragraph("HealthGuard Uganda proposes a resilient offline-first digital health platform to counter health misinformation, streamline household-level triage, and strengthen central coordination across Uganda's rural and semi-urban health system. The solution fuses:")
doc.add_paragraph("1. An offline-first Expo React Native client with on-device edge AI for rumor classification.", style='List Bullet')
doc.add_paragraph("2. A central Express/Prisma coordination backend that synchronizes data when connectivity returns.", style='List Bullet')
doc.add_paragraph("3. A hospital landing portal for regional public-facing information and scheduling.", style='List Bullet')
doc.add_paragraph("Rationale and evidence: offline-first architectures maintain data collection and care processes in low-connectivity settings; edge AI enables real-time misinformation screening without network access; centralized dashboards and RBAC support outbreak and rumor surveillance at district/national scales. The platform targets measurable outcomes including reductions in non-critical hospital visits, improved household maternal tracking, and enhanced district visibility into rumor trends and service utilization.")

# Problem Statement
doc.add_heading('Problem Statement', level=1)
doc.add_paragraph("Uganda's community health system faces four interlocking bottlenecks that HealthGuard is designed to alleviate:")

doc.add_heading('1) Proliferation of health misinformation and infodemics', level=2)
doc.add_paragraph("The infodemic undermines timely care and vaccine uptake; locally contextualized counter-messaging is essential. Integrated, on-device screening plus a central repository of verified responses aligns with best-practice infodemic management in low-resource settings.")

doc.add_heading('2) Intermittent connectivity', level=2)
doc.add_paragraph("Rural CHWs operate with limited data connectivity, diminishing cloud-centric tools. Offline-first architectures with local data stores and deferred synchronization have demonstrated viability for continuous data collection and care delivery.")

doc.add_heading('3) Strained healthcare infrastructure', level=2)
doc.add_paragraph("Primary centers manage high patient volumes; standardized community-level triage and maternal health tracking can reduce unnecessary ER visits and improve ANC adherence.")

doc.add_heading('4) Lack of central coordination', level=2)
doc.add_paragraph("District and national health offices require real-time visibility into rumor trends and local service utilization to mount timely responses; dashboards and geo-mapped data are central to this capability.")

# Solution Overview
doc.add_heading('Solution Overview', level=1)
doc.add_paragraph("HealthGuard Uganda deploys a tri-layer architecture designed for resilience, speed, and scale:")
doc.add_paragraph("Offline-First Edge Client: An Expo React Native application using expo-sqlite to enable complete offline operation for case screening, patient detail logging, maternal ANC tracking, and access to a verified knowledge base.")
doc.add_paragraph("Edge AI Misinformation Classifier: A lightweight on-device logistic regression model that screens incoming claims and stories for misinformation likelihood, enabling immediate triage and counseling without network access.")
doc.add_paragraph("National Backend Coordination Server: Express.js API with Prisma ORM, backed by PostgreSQL/SQLite. Services include central sync, real-time dashboards, and RBAC.")
doc.add_paragraph("Hospital Public Portal: A responsive static site with hospital departments, referral pathways, and emergency contacts.")

# Development Phase & Implementation Timeline
doc.add_heading('Development Phase & Implementation Timeline', level=1)
doc.add_paragraph("Phase 1: Foundation & Baseline (Months 1–2)", style='List Bullet')
doc.add_paragraph("Phase 2: Local Intelligence & Offline Core (Months 3–4)", style='List Bullet')
doc.add_paragraph("Phase 3: Backend & Sync Synchronization (Months 5–6)", style='List Bullet')
doc.add_paragraph("Phase 4: Pilot & Field Deployment (Month 7+)", style='List Bullet')

# Metrics, Evaluation, and Impact
doc.add_heading('Metrics, Evaluation, and Impact', level=1)
doc.add_paragraph("Public Health Advocacy: Proportion of locally verified claims matched at point of care; rate and speed of rumor flagging.")
doc.add_paragraph("Operational Efficiency: Reduction in client check-in times; ANC adherence rates; completion rate of offline triage workflows.")
doc.add_paragraph("National Coordination: Lead-time to address district outbreaks; dashboard completeness.")
doc.add_paragraph("Safety & Privacy: Maturity of RBAC implementation; encryption; data retention policies.")
doc.add_paragraph("Equity & Access: Reach via planned USSD/SMS fallback; multilingual content coverage.")
doc.add_paragraph("Sustainability: Modular, open-source-friendly architecture; local capacity-building.")

# Future Scalability & Next Steps
doc.add_heading('Future Scalability & Next Steps', level=1)
doc.add_paragraph("Regional Language Localization: Luganda, Runyankole, Acholi, etc., to broaden accessibility.")
doc.add_paragraph("USSD/SMS Fallback: Extend core capabilities to feature phones via USSD gateways.")
doc.add_paragraph("Advanced ML & LLM Integration: Introduce deeper local models as hardware allows.")
doc.add_paragraph("National Rollout Strategy: Phased expansion with MoH partnerships.")

# Annexes
doc.add_heading('Annex A: Data Model (High-Level)', level=1)
doc.add_paragraph("Entities: User, Role, Household, Patient, Case, TriageLog, MaternalRecord, KnowledgeAsset, Claim, VerificationRecord, SyncPacket, AuditLog, HospitalSitePage, Appointment, District, HealthEvent.")
doc.add_paragraph("Relationships: Users assign Roles; Households contain Patients; Patients have Cases, etc.")

doc.add_heading('Annex B: Risk Register', level=1)
doc.add_paragraph("Connectivity Variability: Mitigation – robust offline storage, scheduled sync.")
doc.add_paragraph("Data Privacy/Regulatory Compliance: Mitigation – RBAC, encryption, data minimization.")
doc.add_paragraph("Model Drift/Edge AI Performance: Mitigation – versioned models, monitoring.")

doc.add_heading('Annex C: Budget Outline', level=1)
doc.add_paragraph("Personnel: Product manager, software engineers, UX designer.")
doc.add_paragraph("Hardware: Mobile devices for pilots, offline storage, server infrastructure.")
doc.add_paragraph("Training & capacity-building: CHW training sessions.")
doc.add_paragraph("Contingency: 10–15%.")

doc.add_heading('Annex D: Glossary', level=1)
doc.add_paragraph("Offline-First: Application functions fully without network access.")
doc.add_paragraph("Edge AI: On-device machine learning inference.")
doc.add_paragraph("RBAC: Role-Based Access Control.")
doc.add_paragraph("Bi-Directional Sync: Two-way data synchronization between local and central stores.")

# Save the document
doc.save('HealthGuard_Uganda_Proposal.docx')
print("Document saved to HealthGuard_Uganda_Proposal.docx")
